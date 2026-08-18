from pywinauto import Application
from openpyxl import load_workbook
import time

# =============================================
# ERP 근태계획 자동입력 스크립트
# =============================================

def get_main_window():
    """ERP 창 연결"""
    app = Application(backend="uia").connect(title="존슨일렉트릭오퍼레이션스")
    return app.top_window()

def click_button(win, btn_name):
    """툴바 버튼 클릭 (신규/저장/출력)"""
    toolbar = win.child_window(auto_id="tbToolBar", control_type="ToolBar")
    btn = toolbar.child_window(title=btn_name, control_type="Text")
    btn.click_input()
    time.sleep(1.5)  # 0.8 → 1.5로 증가

def set_edit_field(win, auto_id, value):
    """텍스트 입력 필드에 값 입력"""
    custom = win.child_window(auto_id=auto_id, control_type="Custom")
    edit = custom.child_window(auto_id="textBox1", control_type="Edit")
    edit.set_focus()
    time.sleep(0.3)
    # triple_click 대신 Ctrl+A로 전체선택 후 입력
    edit.type_keys("^a")
    edit.type_keys(str(value), with_spaces=True)
    time.sleep(0.3)

def set_combo_field(win, auto_id, value):
    """드롭다운(ComboBox)에서 항목 선택"""
    custom = win.child_window(auto_id=auto_id, control_type="Custom")
    combo = custom.child_window(control_type="ComboBox")
    combo.set_focus()
    time.sleep(0.3)
    # 드롭다운 먼저 열기
    combo.click_input()
    time.sleep(0.5)
    # ListItem 클릭
    item = combo.child_window(title=value, control_type="ListItem")
    item.click_input()
    time.sleep(0.3)

def set_date_field(win, auto_id, date_str):
    """날짜 필드 입력 (형식: 2026-04-08)"""
    custom = win.child_window(auto_id=auto_id, control_type="Custom")
    edit = custom.child_window(auto_id="EditDate1", control_type="Edit")
    edit.set_focus()
    time.sleep(0.3)
    edit.type_keys("^a")
    # 날짜 형식 변환: 2026-04-08 → 20260408
    date_clean = str(date_str).replace("-", "")
    edit.type_keys(date_clean)
    time.sleep(0.3)

def set_time_field(win, auto_id, time_str):
    """시간 필드 입력 (형식: 09:00)"""
    custom = win.child_window(auto_id=auto_id, control_type="Custom")
    edit = custom.child_window(auto_id="textBox1", control_type="Edit")
    edit.set_focus()
    time.sleep(0.3)
    edit.type_keys("^a")
    # 시간 형식 변환: 09:00 → 0900
    time_clean = str(time_str).replace(":", "")
    edit.type_keys(time_clean)
    time.sleep(0.3)

def input_one_row(win, row):
    """한 행 데이터를 ERP에 입력"""
    print(f"  → 신규 클릭")
    click_button(win, "신규")

    print(f"  → 사원명: {row['사원명']}")
    set_edit_field(win, "txtEmpName", row['사원명'])

    print(f"  → 근태: {row['근태']}")
    set_combo_field(win, "cmbUMItemTypeName", row['근태'])

    print(f"  → 사유: {row['사유']}")
    set_edit_field(win, "txtDescription", row['사유'])

    print(f"  → 근태시작일: {row['시작일']}")
    set_date_field(win, "datStartDate", row['시작일'])

    print(f"  → 근태시작시간: {row['시작시간']}")
    set_time_field(win, "mskStartTime", row['시작시간'])

    print(f"  → 근태종료일: {row['종료일']}")
    set_date_field(win, "datEndDate", row['종료일'])

    print(f"  → 근태종료시간: {row['종료시간']}")
    set_time_field(win, "mskEndTime", row['종료시간'])

    print(f"  → 배차여부: {row['배차여부']}")
    set_combo_field(win, "cmbUMCarUseName", row['배차여부'])

    print(f"  → 연락처: {row['연락처']}")
    set_edit_field(win, "txtCellPhone", row['연락처'])

    print(f"  → 저장")
    click_button(win, "저장")

    print(f"  → 출력")
    win.child_window(auto_id="Print", control_type="Button").click_input()
    time.sleep(1)

def run():
    """메인 실행 함수"""
    # Excel 파일 읽기
    wb = load_workbook("근태입력.xlsx")
    ws = wb.active

    # 헤더 행 읽기 (1행)
    headers = [cell.value for cell in ws[1]]
    print(f"컬럼 확인: {headers}")

    # ERP 창 연결
    win = get_main_window()
    print("✅ ERP 연결 성공\n")

    # 결과 로그
    success_list = []
    fail_list = []

    # 2행부터 데이터 처리
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):  # 빈 행 스킵
            continue

        row_data = dict(zip(headers, row))
        emp_name = row_data.get('사원명', '')
        print(f"\n[{row_idx}행] {emp_name} 입력 시작...")

        try:
            input_one_row(win, row_data)
            success_list.append(emp_name)
            print(f"✅ {emp_name} 완료!")

        except Exception as e:
            fail_list.append({"행": row_idx, "사원명": emp_name, "오류": str(e)})
            print(f"❌ {emp_name} 실패: {e}")

        time.sleep(1)  # 다음 행 전 대기

    # 결과 요약
    print("\n" + "="*40)
    print(f"✅ 성공: {len(success_list)}건")
    print(f"❌ 실패: {len(fail_list)}건")
    if fail_list:
        print("\n실패 목록:")
        for f in fail_list:
            print(f"  - {f['행']}행 {f['사원명']}: {f['오류']}")

        # 실패 건 Excel 저장
        from openpyxl import Workbook
        wb_fail = Workbook()
        ws_fail = wb_fail.active
        ws_fail.append(["행", "사원명", "오류"])
        for f in fail_list:
            ws_fail.append([f["행"], f["사원명"], f["오류"]])
        wb_fail.save("실패_로그.xlsx")
        print("→ 실패_로그.xlsx 저장됨")

if __name__ == "__main__":
    run()