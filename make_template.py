from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "근태입력"

# 헤더 (스크립트의 컬럼명과 반드시 일치)
headers = ["사원명", "근태", "사유", "시작일", "시작시간", "종료일", "종료시간", "배차여부", "연락처"]
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center")

# 근태 드롭다운
dv_geuntae = DataValidation(
    type="list",
    formula1='"사용외출,공용외출,출장,조퇴,월차휴가,년차휴가,생차휴가,유급휴가,무급휴가,훈련,사외교육,휴직,조합활동,유선결근,무단결근,공상,산업재해,오전반차,오후반차"',
    showDropDown=False
)
ws.add_data_validation(dv_geuntae)
dv_geuntae.sqref = "B2:B1000"

# 배차여부 드롭다운
dv_baecha = DataValidation(type="list", formula1='"예,아니오"', showDropDown=False)
ws.add_data_validation(dv_baecha)
dv_baecha.sqref = "H2:H1000"

# 샘플 데이터
sample = ["이진용", "사용외출", "병원방문", "2026-04-08", "09:00", "2026-04-08", "12:00", "아니오", "010-1234-5678"]
ws.append(sample)

# 열 너비 조정
widths = [12, 12, 20, 14, 10, 14, 10, 10, 16]
for col, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(1, col).column_letter].width = w

wb.save("근태입력.xlsx")
print("✅ 근태입력.xlsx 생성 완료!")